# -*- coding: utf-8 -*-
"""주간 정보보호 동향 보고서 (2026-09-28 ~ 2026-10-05) .docx 생성 스크립트"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/user/Study/reports/주간_정보보호_동향_보고서_2026-10-05.docx"
FONT = "맑은 고딕"
NAVY = "1F3864"

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.8)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def run(p, text, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(text="", bold=False, size=None, color=None, align=None, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    if text:
        run(p, text, bold, size, color)
    return p


def h1(text):
    p = para(after=4)
    p.paragraph_format.space_before = Pt(10)
    run(p, text, True, 14, NAVY)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", NAVY)):
        bt.set(qn("w:" + k), v)
    b.append(bt)
    pPr.append(b)


def h2(text, color=NAVY):
    p = para(after=2)
    p.paragraph_format.space_before = Pt(6)
    run(p, text, True, 11, color)


def h3(text):
    p = para(after=1)
    p.paragraph_format.space_before = Pt(4)
    run(p, text, True, 10, "2F5496")


def bullet(text, boldlead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    if boldlead:
        run(p, boldlead, True)
    run(p, text)
    return p


def box(lines, fill="EAF1FB", title=None):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, fill)
    c.paragraphs[0].paragraph_format.space_after = Pt(2)
    first = True
    if title:
        run(c.paragraphs[0], title, True, 10.5, NAVY)
        first = False
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        run(p, ln)
    para(after=2)


SEV = {"C": ("🔴 CRITICAL", "F8D7DA"), "H": ("🟠 HIGH", "FDE9D0"), "M": ("🟡 MEDIUM", "FFF6CC")}


def table(headers, rows, widths=None, sevcol=None, size=8.5, head=NAVY):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, head)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(c.paragraphs[0], h, True, size, "FFFFFF")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            run(cells[i].paragraphs[0], str(v), size=size)
            if sevcol is not None and i == sevcol:
                key = "C" if "CRITICAL" in v or "위기" in v or "1순위" in v and "🔴" in v else None
                if "🔴" in v:
                    shade(cells[i], "F8D7DA")
                elif "🟠" in v:
                    shade(cells[i], "FDE9D0")
                elif "🟡" in v:
                    shade(cells[i], "FFF6CC")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    para(after=3)
    return t


# ------------------------------------------------------------------ 헤더
para("주간 정보보호 동향 보고서", True, 22, NAVY, WD_ALIGN_PARAGRAPH.LEFT, 2)
para("2026-09-28 ~ 2026-10-05 · 침해사고대응팀", False, 11, "595959", after=2)
para("보고 대상: CEO · CIO · CISO · CPO · 정보보호부장　|　작성: 정보보호부 침해사고대응팀　|　분류: 사내 한정 (Confidential)",
     size=8.5, color="595959", after=2)
para("수집 기준: KST 기준 보고서 작성일(2026-10-05) 포함 직전 7일(2026-09-28 ~ 2026-10-05). "
     "원문 사이트 다수가 실행 환경의 네트워크 정책으로 직접 열람이 불가하여, 웹 검색 결과에 표시된 기사·공지 요약과 "
     "복수 출처 대조로 사실을 확인하였다(세부는 자체검증 참조).", size=8, color="7F7F7F", after=4)

# ------------------------------------------------------------------ SECTION 1
h1("SECTION 1 — EXECUTIVE SUMMARY")
h2("금주 보안 기상도")
table(["등급", "사유"], [
    ["🔴 위기", "신한은행 등 금융권 7개사 연쇄 해킹(신한은행 25,729명 연소득·대출한도 등 신용정보 유출, AI 자율 침투 도구 흔적 확인)과 "
               "Citrix NetScaler(CVE-2026-88771/88772)·FortiMail(CVE-2026-104286) 인증 없는 원격 악용(CISA KEV 등재)이 동시에 확인되었다."],
    ["🟠 경계", "대출모집인 조회 등 외부 노출 API의 권한검증 미흡 사례(무신사·강남언니·위버스 포함)와 AI 자율 침투 도구 기반 광범위 스캐닝, "
               "Citrix SAML 취약점(CVE-2026-88779)·Zammad(콜센터·헬프데스크 도구) KEV 등재, Warlock 랜섬웨어의 SharePoint 침투가 확인되었다."],
    ["🟡 주의", "금융위원회 긴급 상황대응 회의·보안점검 지시 및 대통령 지시(제도 강화 예고), KillSec 랜섬웨어 일제 단속, GitLab AI Gateway 치명적 RCE 패치, ASEC 다크웹 동향."],
], widths=[2.2, 15.2], sevcol=0, size=9)

h2("핵심 위협 요약")
box([
    "글로벌에서는 외부에 노출된 보안 장비·메일 게이트웨이가 주된 침투 경로로 확인되었다. Citrix NetScaler ADC·Gateway의 인증 없는 원격 명령 실행 취약점"
    "(CVE-2026-88771·88772, CVSS v4.0 9.5)이 패치 공지일(2026-09-27) CISA KEV에 등재되었고, FortiMail(CVE-2026-104286, CVSS 9.8)은 "
    "패치가 제공되지 않은 상태에서 악용이 확인되어 CISA KEV에 등재되었다(2026-10-01). "
    "Symantec은 Warlock 랜섬웨어 조직이 1년 전 공개된 SharePoint 취약점(ToolShell)을 여전히 초기 침투에 사용하고 있다고 공개하였다(2026-10-02).",
    "국내에서는 2026-09-30~10-02 신한은행·KB국민은행·하나은행·BNK부산은행·예가람저축은행·웰컴저축은행·현대캐피탈 등 금융 7개사에서 개인(신용)정보 유출이 잇따라 공개되었다. "
    "신한은행은 고객 25,729명(고객명·전화번호·연소득·대출 산출한도 등)이며, 금융보안원의 공격 IP 역추적에서 중국어 기반 오픈소스 AI 자율 침투 도구 ARTEX 구동 정황이 확인되었다. "
    "무신사·강남언니·위버스 등 비금융 사고도 API 권한검증 미흡이라는 공통점이 보도되었다.",
    "금융위원회는 2026-10-02 긴급 상황대응 회의에서 은행·카드업권에 외부 노출 전산시스템 전반의 보안점검을 주문하였고, 2026-10-04 금융회사 CEO를 소집하였으며, 대통령이 철저한 조사와 대책 마련을 지시하였다.",
    "당사 관점의 가장 긴급한 조치 사항은 대출모집인·제휴·모바일 채널 등 외부 노출 API의 인증·권한검증 전수 점검과 Citrix NetScaler·FortiMail 보유 여부 확인 및 패치·침해 흔적 점검이다.",
], title="이번 주 핵심 위협")

h2("⚡ 긴급조치 Top 3")
for i, t in enumerate([
    "외부 노출 API·조회 페이지(대출모집인·제휴사·앱 연계) 인증·객체 단위 권한검증·호출량 제한 전수 점검 — 정보보호부 / 채널·제휴 시스템 운영팀",
    "Citrix NetScaler(CVE-2026-88771·88772·88779) 보유 자산 식별, 벤더 패치(CTX697096) 적용 및 침해 흔적 점검 — 인프라보안팀 / VPN·SSO 게이트웨이",
    "FortiMail(7.2~8.0 계열) 보유 확인, IBE 기능 비활성화 및 관리 인터페이스 인터넷 차단 — 메일보안 담당 / 메일 게이트웨이",
], 1):
    p = para(after=1)
    run(p, f"{i}. ", True, color="C00000")
    run(p, t)

# ------------------------------------------------------------------ SECTION 2
h1("SECTION 2 — 사건 및 위협 현황 리스트")
para("수집 13건 전체를 심각도 내림차순으로 기재한다. 날짜는 발생·공개일(KST 기준 보도·공지일)이다.", size=8.5, color="595959")
S2 = [
    ["🔴 CRITICAL", "개인정보침해", "신한은행 대출모집인 조회 서비스 침해, 고객 25,729명 신용정보 유출", "신한은행(국내)", "2026-10-01", "✅", "높음"],
    ["🔴 CRITICAL", "취약점 악용", "Citrix NetScaler ADC·Gateway 제로데이 RCE (CVE-2026-88771/88772, CVSS v4.0 9.5)", "NetScaler 이용 전 세계 조직", "2026-09-27", "✅", "높음"],
    ["🔴 CRITICAL", "취약점 악용", "FortiMail 경로 순회·NULL 바이트 취약점 악용 (CVE-2026-104286, CVSS 9.8, 패치 미제공)", "FortiMail 이용 조직", "2026-10-01", "✅", "높음"],
    ["🟠 HIGH", "개인정보침해", "예가람저축은행(약 4만 명)·웰컴저축은행(법인 최대 2,200건)·현대캐피탈(모집인 146명)·KB국민(119명)·하나(89명)·BNK부산(11명) 연쇄 유출", "국내 은행·저축은행·캐피탈", "2026-09-30 ~ 10-02", "✅", "높음"],
    ["🟠 HIGH", "APT·해킹", "금융보안원, AI 자율 침투 도구(ARTEX) 기반 금융권 광범위 공격 진행 권고", "국내 금융회사", "2026-10-02", "✅", "높음"],
    ["🟠 HIGH", "개인정보침해", "무신사(29CM)·강남언니·위버스 API 권한검증 미흡 기반 개인정보 유출", "국내 이커머스·의료플랫폼·팬덤플랫폼", "2026-09-28(보도)", "✅", "중간"],
    ["🟠 HIGH", "취약점 악용", "Citrix NetScaler SAML 메모리 오버플로우 (CVE-2026-88779, CVSS 8.7), 패치 적용 장비 대상 악용", "NetScaler(SAML SP/IdP 구성)", "2026-10-04", "✅", "높음"],
    ["🟠 HIGH", "취약점 악용", "Zammad 세션 고정·권한 상승 (CVE-2026-102489/102490) CISA KEV 등재", "Zammad(헬프데스크) 이용 조직", "2026-10-02", "✅", "중간"],
    ["🟠 HIGH", "랜섬웨어", "Warlock 랜섬웨어, SharePoint 취약점으로 수도·통신·지자체·대학 침해 (Symantec)", "중남미·유럽·아프리카 조직", "2026-10-02", "✅", "중간"],
    ["🟡 MEDIUM", "취약점 악용", "GitLab AI Gateway 명령 실행 취약점 (CVE-2026-90970, CVSS 9.9, 인증 필요) 패치", "자체 호스팅 GitLab AI Gateway", "2026-10 (1주차)", "✅", "중간"],
    ["🟡 MEDIUM", "랜섬웨어", "KillSec 랜섬웨어 조직 일제 단속(Operation KillSwitch), 체포 3명", "글로벌(법집행)", "2026-09-30", "✅", "낮음"],
    ["🟡 MEDIUM", "정책·제도", "금융위 긴급 상황대응 회의·외부 노출 시스템 보안점검 지시, CEO 소집, 대통령 지시", "국내 전 금융권", "2026-10-02 ~ 10-04", "✅", "높음"],
    ["🟡 MEDIUM", "랜섬웨어", "ASEC 다크웹 동향 1주차: 국내 종합병원 유출 주장, Qilin 일본 물류사 공격 등", "국내 병원·일본 물류사", "2026-10 (1주차)", "⚠️", "낮음"],
]
table(["심각도", "위협유형", "사건명", "피해 주체", "발생·공개일", "교차검증", "금융권 관련성"], S2,
      widths=[2.0, 1.8, 6.4, 3.0, 2.0, 1.1, 1.4], sevcol=0, size=8)

h3("Tier별 탐색 결과")
table(["Tier", "탐색 결과"], [
    ["Tier 1 공공·정부·감독기관", "금융위원회(보도자료)·금융보안원(권고, 보도 인용)·금융감독원(현장조사, 보도 인용)·CISA(KEV 4건 공지) 확인. "
                              "KISA·개인정보보호위원회·과기정통부·국가사이버안보센터·경찰청·NIST·ENISA: 해당 기간 특이사항 없음(검색에서 기간 내 관련 자료 미확인)."],
    ["Tier 2 보안기업", "AhnLab ASEC(다크웹 주간 동향), Symantec(Warlock), Truesec·watchTowr(취약점 분석) 확인. "
                      "S2W·Mandiant·CrowdStrike·Unit 42·Check Point·Talos·Cloudflare·Akamai·NETSCOUT·SK쉴더스: 해당 기간 특이사항 없음(기간 내 관련 자료 미확인)."],
    ["Tier 3 전문 언론", "BleepingComputer·SecurityWeek·The Hacker News·Dark Reading·보안뉴스·데일리시큐 확인."],
    ["Tier 4 일반 언론", "서울신문·한국일보·파이낸셜뉴스·머니투데이·뉴시스·뉴데일리·뉴스1 등 확인(사건 최초 보도·규모 확인용)."],
], widths=[3.4, 14.0], size=8.5)

# ------------------------------------------------------------------ SECTION 3
h1("SECTION 3 — 위협 유형별 통계")
table(["위협유형", "건수", "심각도 분포"], [
    ["취약점 악용", "5", "🔴 2 / 🟠 2 / 🟡 1"],
    ["랜섬웨어", "3", "🟠 1 / 🟡 2"],
    ["APT·해킹", "1", "🟠 1"],
    ["개인정보침해", "3", "🔴 1 / 🟠 2"],
    ["정책·제도", "1", "🟡 1"],
    ["**총계**".replace("**", ""), "13", "🔴 3 / 🟠 6 / 🟡 4"],
], widths=[5, 2.5, 9.9], size=9)
para("※ DDoS·공급망·피싱·사회공학: 해당 기간 내 금융권 관련 확인 사건 없음(0건, 표 미표기).", size=8.5, color="595959")

# ------------------------------------------------------------------ SECTION 4
h1("SECTION 4 — 주요 사건 상세 보고")
para("SECTION 2의 13건 중 금융권 영향도·심각도가 높은 6건을 상세 분석한다. "
     "※ MITRE ATT&CK ID는 보도·분석 내용에 근거한 분류이며, 공식 조사 결과로 확정된 것은 아니다.", size=8.5, color="595959")


def _render_event(title, color, who, overview, meta, cond, sources, cross, risk, impact, action):
    h2(title, color)
    h3("사건 개요")
    for k, v in zip(["누가", "언제", "어디서", "무엇을", "어떻게", "얼마나"], overview):
        bullet(" " + v, f"{k}:")
    h3("메타데이터")
    para("[필수]", True, 9)
    for k, v in meta:
        bullet(" " + v, f"{k}:")
    if cond:
        para("[조건부]", True, 9)
        for k, v in cond:
            bullet(" " + v, f"{k}:")
    h3("출처 및 교차확인")
    for i, s in enumerate(sources, 1):
        bullet(" " + s, f"출처 {i}:")
    bullet(" " + cross, "교차확인:")
    h3("💳 금융권 시사점")
    bullet(" " + risk, "[리스크]")
    bullet(" " + impact, "[영향]")
    bullet(" " + action, "[조치]")


EV = []


def event(*a):
    EV.append(a)


LAW = ("개인정보 보호법, 신용정보의 이용 및 보호에 관한 법률(신용정보법), 전자금융거래법 (조문 확인 필요)")

event("🔴 [개인정보침해] 신한은행 — 대출모집인 조회 서비스 등 외부 노출 서비스 침해, 고객 25,729명 개인(신용)정보 유출", "C00000", None,
      ["신한은행(피해 주체). 공격 주체는 미공개이며 귀속 확인되지 않음(조사 중).",
       "2026-10-01 유출 공개, 금융감독원 은행검사국·IT검사국 현장조사 착수.",
       "신한은행 대출모집인용 조회 서비스 등 외부 노출 서비스(국내).",
       "고객명·전화번호·연소득·대출 산출한도 등 대출 신청 관련 개인(신용)정보. 일부 고객의 주민등록번호 66건, CI(연계정보) 97건 포함 보도.",
       "외부 노출 서비스·API를 통한 조회로 유출된 것으로 보도됨. 금융보안원이 공격 IP를 역추적한 결과 공격 서버에서 AI 자율 침투 도구 ARTEX 흔적이 확인됨(상세 아래 ARTEX 항목). 구체 경로는 조사 중.",
       "고객 25,729명(은행 공개 기준, 이후 금융권 7개사 합산 약 6.6만 명 보도)."],
      [("발생·공개일", "2026-10-01"),
       ("공격 유형", "T1190 외부 노출 애플리케이션 악용(보도 기반 분류, 확정 아님) / 상세 경로 조사 중"),
       ("피해 규모", "25,729명")],
      [("유출 항목", "고객명, 전화번호, 연소득, 대출 산출한도, (일부) 주민등록번호 66건·CI 97건"),
       ("관련 규제", LAW)],
      ["[T4] 서울신문 / [속보] 신한은행 2만5천명 정보 유출…연소득·대출정보까지 털렸다 / 2026-10-01 / https://www.seoul.co.kr/news/economy/2026/10/01/20261001500101",
       "[T4] 한국일보 / 신한은행 해킹에 2만5000명 개인정보 유출… 금감원 긴급 현장조사 / 2026-10-01 / https://www.hankookilbo.com/news/article/A2026100115320000186",
       "[T4] 파이낸셜뉴스 / 금융사 7곳 뚫렸다… 개인 6만6천명·법인 2200건 정보 유출 / 2026-10-04 / https://fnnews.com/news/202610041822434160"],
      "✅ 확인 (3개 소스 / 최고 신뢰도: T4)",
      "신한은행 사례에서 연소득·대출한도 등 신용정보가 포함되어 신용정보법상 신용정보 유출 및 개인정보 보호법상 유출 통지·신고 의무, 전자금융거래법상 안전성 확보 의무 관련 감독 리스크가 발생하였다. 금융감독원이 현장조사 중이므로 당사도 동종 서비스에 대한 감독 점검 대상이 될 수 있다.",
      "당사가 대출모집인·제휴 채널에 조회 기능을 제공하고 해당 기능의 인증 또는 객체 단위 권한검증이 미흡한 경우, 신한은행과 같이 정상 요청 형태의 대량 조회로 연소득·한도 정보가 유출되어 고객 민원·보상·감독 제재 리스크가 발생할 수 있다.",
      "정보보호부는 모집인·제휴 채널 조회 API의 인증·권한검증·호출량 제한 점검을 즉시 수행하고, 채널운영팀은 비정상 대량조회 탐지 룰을 보안관제에 반영한다.")

event("🟠 [개인정보침해] 예가람저축은행·웰컴저축은행·현대캐피탈·KB국민·하나·BNK부산 — 동일 공격자 IP 다수 확인, 금융 7개사 연쇄 유출", "C55A11", None,
      ["신한은행 외 6개 금융사(피해 주체). 파이낸셜뉴스·SBS Biz는 여러 침해 사고에서 동일 공격자 IP가 확인되었다고 보도. 공격 주체는 미공개(귀속 확인되지 않음).",
       "예가람저축은행 2026-09-30 서버 접근 정황 확인, 현대캐피탈 2026-10-02 확인, KB국민·하나은행 2026-10-02 공개.",
       "예가람저축은행 고객정보 서버, 현대캐피탈 주택대출 모집인 조회페이지, KB국민·하나·BNK부산·웰컴 각 서비스(국내).",
       "예가람: 성명·생년월일·연락처. 현대캐피탈: 모집인 이름·휴대전화번호·이메일·내부 모집인번호·주민등록번호. KB국민: 고객명·전화번호·주소·암호화된 주민번호. 하나: 주민번호·이름·주소·전화번호.",
       "해외 IP를 통한 공격으로 현대캐피탈이 확인하였으며(AI 활용 주장은 회사 발표 기준), 개별 기법은 조사 중.",
       "예가람 약 4만 명(추정), 웰컴 법인 최대 2,200건, 현대캐피탈 모집인 146명, KB국민 119명(고객 99·임직원 20), 하나 89명, BNK 외주 개발직원 11명. 신한 포함 개인 약 6.6만 명."],
      [("발생·공개일", "2026-09-30 ~ 2026-10-02"),
       ("공격 유형", "조사 중 (외부 노출 서비스 대상 공격, T1190 후보)"),
       ("피해 규모", "개인 약 6.6만 명 + 법인정보 최대 2,200건")],
      [("유출 항목", "성명, 생년월일, 연락처, 주소, 주민등록번호(일부 암호화), 이메일, 내부 모집인번호"),
       ("관련 규제", LAW)],
      ["[T4] 파이낸셜뉴스 / 금융사 7곳 뚫렸다… 개인 6만6천명·법인 2200건 정보 유출 / 2026-10-04 / https://fnnews.com/news/202610041822434160",
       "[T3] 보안뉴스 / [은행이 뚫린다] 현대캐피탈 해킹 공격... 주택대출 모집인 146명 정보 유출 / 2026-10 (일자 미표기) / https://www.boannews.com/news/articleView.html?idxno=146187",
       "[T4] 서울경제 / 현대캐피탈·예가람저축은행도 뚫렸다…금융권, 개인정보 잇단 유출 / 2026-10 (일자 미표기) / https://m.sedaily.com/amparticle/20097944"],
      "✅ 확인 (3개 소스 / 최고 신뢰도: T3)",
      "캐피탈·저축은행·은행 등 업권 구분 없이 유출이 확인되어, 여신금융업(카드·캐피탈) 전반이 동일 유형 점검 대상이 될 수 있다. 개인정보 보호법·신용정보법·전자금융거래법 관련 신고·통지 및 감독기관 점검 부담이 따른다(조문 확인 필요).",
      "당사가 모집인·제휴 법인·외주 개발인력 정보를 조회·관리하는 페이지를 외부에 노출하고 있고 해당 페이지에 인증 미흡이 있는 경우, 현대캐피탈 사례처럼 임직원·모집인 주민등록번호까지 유출될 수 있다.",
      "정보보호부는 모집인·외주 인력 정보 조회 페이지의 비공개 항목 노출과 주민번호 응답 필드를 점검하고, 개인정보보호책임부서는 모집인 등록정보 최소 수집·마스킹을 적용한다.")

event("🟠 [APT·해킹] 금융권 — AI 자율 침투 도구(ARTEX) 기반 광범위 공격 진행, 금융보안원 권고", "C55A11", None,
      ["공격 주체 미공개(귀속 확인되지 않음). 보도 인용: 금융보안원이 분석·권고, 금융위원회가 긴급 대응회의 개최.",
       "2026-10-02 금융보안원 권고 및 금융위원회 긴급 상황대응 회의.",
       "신한은행 공격 서버 및 동일 IP가 확인된 금융 7개사(국내).",
       "금융회사 외부 노출 웹서비스·API. 금융보안원 분석에서 공격 서버 HTML 제목에 ARTEX 자율 침투 콘솔을 뜻하는 중국어 문자열 확인.",
       "ARTEX는 LLM 기반 중국어권 오픈소스 자율 침투 테스트 시스템으로 정보 수집·취약점 탐색·공격 경로 수립·검증을 자동화하는 도구로 설명됨. 데일리시큐는 API 권한검증 허점을 노린 정상 요청 모방 조회를 공통 패턴으로 분석.",
       "신한은행 로그·공격 IP 기반 확인. 전체 규모는 조사 중(금융 7개사 연관 보도)."],
      [("발생·공개일", "2026-10-02"),
       ("공격 유형", "T1595.002 취약점 스캐닝, T1190 외부 노출 애플리케이션 악용(보도 기반 분류)"),
       ("피해 규모", "조사 중 (금융 7개사 연관)")],
      [("관련 규제", "전자금융감독규정 등 (조문 확인 필요)")],
      ["[T1] 금융위원회 / 금융권 「긴급 상황대응 회의」 개최 / 2026-10-02 / https://www.fsc.go.kr/no010101/87869",
       "[T3] 데일리시큐 / [이슈 진단] ARTEX로 금융권 광범위 공격, 무신사·강남언니·신한은행까지…AI가 파고든 ‘API 권한검증’ 허점 / 2026-10 (일자 미표기) / https://www.dailysecu.com/news/articleView.html?idxno=208718",
       "[T4] 뉴데일리 / 신한은행 공격 서버에 'AI 침투도구' 흔적 … 금융권 '기계 해킹' 현실화되나 / 2026-10-02 / https://biz.newdaily.co.kr/site/data/html/2026/10/02/2026100200231.html"],
      "✅ 확인 (3개 소스 / 최고 신뢰도: T1)",
      "금융보안원이 광범위 공격 진행을 권고한 상태에서 당사가 관련 위협정보 반영·점검을 지연하면, 감독기관의 점검 요구(체크리스트 회신) 시 이행 미흡으로 평가될 리스크가 있다. 전자금융감독규정상 의무는 조문 확인 필요.",
      "당사 외부 노출 서비스에 권한검증 누락 경로가 존재하는 경우, 자율 침투 도구가 스캐닝·경로 수립을 자동화하므로 사람이 수행하는 공격 대비 탐지 이전에 대량 조회가 완료될 수 있다.",
      "보안관제는 금융보안원 공유 공격 IP·접근 패턴을 WAF·방화벽 차단 및 탐지 룰에 즉시 반영하고, 정보보호부는 호출량·속도 제한과 이상 조회 알림 임계치를 재설정한다.")

event("🟠 [개인정보침해] 무신사(29CM)·강남언니·위버스 — API 권한검증 미흡 악용 개인정보 유출", "C55A11", None,
      ["무신사(29CM), 강남언니(힐링페이퍼), 위버스(피해 주체). 공격 주체 미공개(귀속 확인되지 않음).",
       "보도일 2026-09-28(개별 공개일은 사건별 상이, 조사 중).",
       "회원 주문정보 조회 API(무신사), 웹서비스 상담내역 조회 API(강남언니), 위버스(국내).",
       "회원 주문정보 및 상담내역 등. 강남언니는 시술내역·사진·이름 포함 약 22만 명으로 보도.",
       "API에 정상 요청을 모방한 뒤 조회 권한·호출 횟수 제한 미흡을 악용하는 방식(데일리시큐 분석).",
       "강남언니 약 22만 명(보도), 무신사·위버스 규모는 조사 중."],
      [("발생·공개일", "2026-09-28 (보도일)"),
       ("공격 유형", "T1190 외부 노출 애플리케이션 악용 (API 인가 우회, 보도 기반 분류)"),
       ("피해 규모", "강남언니 약 22만 명, 그 외 조사 중")],
      [("유출 항목", "주문정보, 상담내역, 시술내역, 사진, 이름"),
       ("관련 규제", "개인정보 보호법 (조문 확인 필요)")],
      ["[T4] 머니투데이 / 무신사·강남언니·위버스…해킹사고 모두 '이곳'이 뚫렸다 / 2026-09-28 / https://www.mt.co.kr/tech/2026/09/28/2026092810440149574",
       "[T3] 데일리시큐 / [이슈 진단] ARTEX로 금융권 광범위 공격…AI가 파고든 ‘API 권한검증’ 허점 / 2026-10 (일자 미표기) / https://www.dailysecu.com/news/articleView.html?idxno=208718"],
      "✅ 확인 (2개 소스 / 최고 신뢰도: T3)",
      "비금융 플랫폼 사고이나 금융권 사고와 동일한 API 권한검증 취약점 유형이 확인되어, 당사 마이데이터·제휴몰·간편결제 연계 채널이 같은 유형으로 평가될 수 있다.",
      "당사 제휴 가맹점·플랫폼 API가 사용자 식별자 변경만으로 타인 정보를 조회할 수 있는 구조(OWASP API Security BOLA, 객체 수준 권한 검증 실패)인 경우, 제휴사 사고가 당사 회원정보 유출로 이어질 수 있다.",
      "정보보호부는 OWASP API Top 10 기준 BOLA·과다 데이터 노출 점검을 제휴 연계 API에 우선 수행하고, 제휴 관리부서는 제휴사 보안점검 결과 제출을 요청한다.")

event("🔴 [취약점 악용] Citrix NetScaler ADC·Gateway — 인증 없는 RCE 제로데이(CVE-2026-88771/88772), 전 세계 악용", "C00000", None,
      ["Citrix(NetScaler 제조사) 및 NetScaler 이용 조직(피해 주체). 악용 주체는 미공개.",
       "벤더 패치 번들 CTX697096 및 CISA KEV 등재 2026-09-27. 이후 CVE-2026-88779가 2026-10-04 KEV 등재.",
       "NetScaler ADC·Gateway(VPN·SSO·로드밸런서) 전 세계.",
       "CVE-2026-88771은 입력 검증 오류로 인증 없는 임의 명령 실행이 가능하며 기본 구성에서도 영향. CISA는 전 세계 악용을 확인.",
       "인증 없는 원격 명령 실행. TechTimes는 패치 후에도 백도어가 제거되지 않는다고 보도. SecurityWeek는 패치 적용 며칠 전 장비에서도 악용이 확인되었다고 보도.",
       "악용 건수·피해 조직 수는 미공개(조사 중)."],
      [("발생·공개일", "2026-09-27"),
       ("공격 유형", "T1190 외부 노출 애플리케이션 악용, T1505.003 웹셸(백도어 보도 기반)"),
       ("피해 규모", "조사 중 (전 세계 악용)")],
      [("CVE", "CVE-2026-88771·CVE-2026-88772 (CVSS v4.0 9.5 / Critical, 88771 기준), 관련 CVE-2026-88779 (CVSS 8.7 / High, SAML 메모리 오버플로우)"),
       ("영향 버전", "NetScaler ADC·Gateway(세부 버전 범위: 벤더 공지 CTX697096 확인 필요)"),
       ("CISA KEV 등재 여부", "예 (88771·88772: 2026-09-27, 88779: 2026-10-04)"),
       ("패치 기한", "88771·88772: 2026-09-30(CISA 3일 기한, 2차 출처 기준), 88779: 2026-10-07"),
       ("관련 규제", "전자금융감독규정, 개인정보 보호법 안전조치 의무 (조문 확인 필요)")],
      ["[T1] CISA / Critical Zero-Day Vulnerabilities Exploited in Citrix NetScaler ADC, Gateway / 2026-09-27 / https://www.cisa.gov/news-events/alerts/2026/09/27/critical-zero-day-vulnerabilities-exploited-citrix-netscaler-adc-gateway",
       "[T3] SecurityWeek / Citrix Confirms 2 NetScaler Zero-Days After Admins Pulled the Plug / 2026-09 (일자 미표기) / https://www.securityweek.com/citrix-confirms-2-netscaler-zero-days-after-admins-pulled-the-plug/",
       "[T2] watchTowr / Citrix NetScaler Zero-Day RCE FAQ: CVE-2026-88771 and CVE-2026-88772 / 2026-09 (일자 미표기) / https://watchtowr.com/intelligence/citrix-netscaler-zero-day-vulnerabilities-faq/"],
      "✅ 확인 (3개 소스 / 최고 신뢰도: T1)",
      "KEV 등재와 CVSS 9.5가 확인된 취약점으로, NetScaler가 원격접속·SSO 인증 인프라에 쓰이는 경우 침해 시 인증 체계 전반 영향이 가능하여 전자금융거래법·개인정보 보호법상 안전성 확보 의무 위반 리스크가 있다.",
      "당사가 NetScaler를 VPN·SSO·대외 서비스 로드밸런서로 운영하고 미패치 상태인 경우, 인증 없는 원격 명령 실행으로 내부망 진입과 인증 정보 탈취가 가능하다. 패치 후에도 백도어가 잔존할 수 있다는 보도가 있으므로 패치만으로 종결할 수 없다.",
      "인프라보안팀은 NetScaler 자산 식별 후 CTX697096 패치를 즉시 적용하고 웹셸·비정상 프로세스 침해 흔적을 점검한다. 보안관제는 관리 인터페이스 접근 로그를 확인한다.")

event("🔴 [취약점 악용] Fortinet FortiMail — 경로 순회·NULL 바이트 취약점 악용(CVE-2026-104286), 패치 미제공", "C00000", None,
      ["Fortinet(제조사) 및 FortiMail 이용 조직. 악용 주체 미공개.",
       "Fortinet 권고 FG-IR-26-175 및 CISA KEV 등재 2026-10-01.",
       "FortiMail 이메일 게이트웨이(HTTP/HTTPS 인터페이스).",
       "인증 없이 조작된 HTTP(S) 요청으로 운영체제에 임의 파일을 쓸 수 있는 취약점.",
       "경로 순회와 NULL 바이트 처리 미흡을 결합한 임의 파일 쓰기. 보도에 따르면 원격 코드 실행·데이터 유출·내부 이동으로 이어질 수 있음.",
       "피해 조직 수는 미공개. 공식 패치는 확인 시점 기준 미제공."],
      [("발생·공개일", "2026-10-01"),
       ("공격 유형", "T1190 외부 노출 애플리케이션 악용, T1105 도구 전송(임의 파일 쓰기 기반 분류)"),
       ("피해 규모", "조사 중")],
      [("CVE", "CVE-2026-104286 (CVSS 9.8 / Critical)"),
       ("영향 버전", "FortiMail 8.0~8.0.1, 7.6~7.6.6, 7.4~7.4.8, 7.2~7.2.9"),
       ("CISA KEV 등재 여부", "예 (2026-10-01)"),
       ("패치 기한", "2026-10-04 (연방기관 기한, 2차 출처 기준). 패치 미제공으로 우회 조치(IBE 비활성화·관리 인터페이스 차단) 권고"),
       ("관련 규제", "전자금융감독규정, 개인정보 보호법 안전조치 의무 (조문 확인 필요)")],
      ["[T2] Truesec / CVE-2026-104286: FortiMail Path Traversal Vulnerability Actively Exploited / 2026-10 (일자 미표기) / https://www.truesec.com/hub/blog/cve-2026-104286-fortimail-path-traversal-vulnerability-actively-exploited",
       "[T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog / 2026-10-01 / https://www.cisa.gov/news-events/alerts/2026/10/01/cisa-adds-one-known-exploited-vulnerability-catalog",
       "[T3] SecurityOnline / FortiMail Path Traversal Flaw CVE-2026-104286 Exploited in the Wild / 2026-10 (일자 미표기) / https://securityonline.info/fortimail-vulnerability-cve-2026-104286/"],
      "✅ 확인 (3개 소스 / 최고 신뢰도: T1)",
      "메일 게이트웨이는 고객 대상 이메일·내부 문서가 경유하는 지점으로 침해 시 개인(신용)정보 유출 가능성이 있어 개인정보 보호법·신용정보법상 유출 통지 리스크를 수반한다.",
      "당사가 영향 버전 FortiMail을 운영하고 IBE 기능 또는 관리 인터페이스가 인터넷에 노출된 경우, 패치가 없는 상태에서 인증 없는 임의 파일 쓰기로 메일 서버가 장악되고 메일 데이터가 유출될 수 있다.",
      "메일보안 담당은 버전 확인 후 IBE 기능 비활성화와 관리 인터페이스 외부 접근 차단을 즉시 적용하고, 보안관제는 FortiMail 파일 변경·비정상 프로세스를 점검한다. 벤더 패치 공지를 상시 확인한다.")

for a in sorted(EV, key=lambda x: 0 if x[0].startswith("🔴") else 1):
    _render_event(*a)

# ------------------------------------------------------------------ SECTION 5
h1("SECTION 5 — 즉시 조치 권고")
table(["우선순위", "조치 항목", "담당 부서/시스템", "완료 기한", "근거 사건"], [
    ["🔴 1순위", "NetScaler 자산 식별·CTX697096 패치 적용·웹셸 등 침해 흔적 점검", "인프라보안팀 / VPN·SSO 게이트웨이", "즉시", "Citrix 제로데이"],
    ["🔴 1순위", "FortiMail 버전 확인, IBE 비활성화·관리 인터페이스 외부 차단, 파일 쓰기 흔적 점검", "메일보안 담당 / 메일 게이트웨이", "즉시", "FortiMail CVE-2026-104286"],
    ["🟠 2순위", "외부 노출 API·조회 페이지(모집인·제휴·앱 연계)의 인증·객체 단위 권한검증·호출량 제한 점검", "정보보호부 / 채널·제휴 시스템", "금주 내", "신한은행·7개사 유출"],
    ["🟠 2순위", "금융보안원 공유 공격 IP·접근 패턴 WAF·방화벽 차단 및 이상 대량조회 탐지 룰 반영", "보안관제 / WAF·SIEM", "금주 내", "ARTEX 권고"],
    ["🟠 2순위", "금융당국 보안점검 체크리스트 자체 점검 및 결과 보고체계 가동(CEO 보고 포함)", "정보보호부·준법감시", "금주 내", "금융위 회의·대통령 지시"],
    ["🟠 2순위", "NetScaler SAML 구성(SP/IdP) 여부 확인, CVE-2026-88779 대응", "인프라보안팀", "금주 내", "Citrix CVE-2026-88779"],
    ["🟡 3순위", "Zammad·GitLab AI Gateway·온프레미스 SharePoint 보유 확인 및 패치·버전 점검", "IT인프라팀 / 개발·협업 시스템", "이달 내", "Zammad·GitLab·Warlock"],
    ["🟡 3순위", "모집인·외주 인력 개인정보 최소 수집·마스킹, 비공개 항목 응답 필드 점검", "개인정보보호부서 / 채널 개발팀", "이달 내", "현대캐피탈·BNK 사례"],
], widths=[2.0, 7.0, 3.8, 1.6, 3.0], sevcol=0, size=8.5)

# ------------------------------------------------------------------ SECTION 6
h1("SECTION 6 — 사고 개별 분석 리포트")
para("경영진 의사결정이 필요한 3건: ① 금융권 연쇄 유출(AI 자율 침투 도구) ② Citrix NetScaler 제로데이 ③ FortiMail 제로데이.", size=8.5, color="595959")


def report(title, s1, s2, s3, biz, strat, srcs):
    h2("[보고] " + title, NAVY)
    h3("1. 사고 개요 및 규모")
    for k, v in zip(["발견 시점", "유출·피해 규모", "피해 항목"], s1):
        bullet(" " + v, k + ":")
    h3("2. 사고 원인 및 경위")
    for k, v in zip(["공격 유형", "발견 배경", "발생 추정 시점"], s2):
        bullet(" " + v, k + ":")
    h3("3. 현재 대응 상황")
    for k, v in zip(["즉시 대응", "보안 강화", "유관기관 신고"], s3):
        bullet(" " + v, k + ":")
    h3("4. 경영진 제언")
    para("비즈니스 리스크", True, 9)
    for k, v in zip(["법적·규제", "평판·고객", "재무·운영"], biz):
        bullet(" " + v, k + ":")
    para("대응 전략", True, 9)
    for k, v in zip(["즉시 (24시간 이내)", "단기 (1~2주 이내)", "중장기 (1분기 이내)"], strat):
        bullet(" " + v, k + ":")
    h3("5. 관련 출처")
    for s in srcs:
        bullet(" " + s)


report("금융권 연쇄 개인정보 유출(AI 자율 침투 도구 흔적) 사고 분석 및 대응 리포트",
       ["2026-09-30 예가람저축은행 서버 접근 정황 확인, 2026-10-01 신한은행 공개(금융감독원 현장조사) — 확인 주체는 각 회사 및 금융당국",
        "개인 약 6.6만 명, 법인정보 최대 2,200건(신한 25,729명, 예가람 약 4만 명 등). 금융 7개사. 당사 해당 여부: 당사 피해 확인된 바 없음(2026-10-05 기준 공개 자료 한정)",
        "고객명·전화번호·연소득·대출 산출한도, 생년월일·주소, 주민등록번호(일부 암호화), CI 97건, 내부 모집인번호 등"],
       ["외부 노출 서비스·API 대상 침해(T1190), 공격 서버에서 AI 자율 침투 도구 ARTEX 흔적 확인. 기술적 원인: API 권한검증·인증 미흡(보도 분석), 관리적 원인: 조사 중",
        "금융보안원의 신한은행 로그·공격 IP 역추적, 금융사 7곳에서 동일 IP 확인(SBS Biz·파이낸셜뉴스 보도)",
        "2026-09-30 이전(예가람 최초 확인일). 일부 사건은 조사 중"],
       ["신한은행 대국민 사과·피해 전액 보상 표명, 금융감독원 은행검사국·IT검사국 현장조사, 금융위 긴급 상황대응 회의(2026-10-02)",
        "금융위가 외부 노출 시스템 보안점검·인증 및 접근통제 점검·공격 IP 등 위협정보 신속 공유를 지시, 체크리스트 제공 및 결과 보고 예정",
        "금융감독원 현장조사 진행, 2026-10-04 금융회사 CEO 소집, 대통령 철저 조사·범부처 대책 지시. 개인정보보호위원회 조사 여부: 확인되지 않음"],
       ["개인정보 보호법·신용정보법·전자금융거래법상 점검·제재 가능성(조문·제재 수준 확인 필요). 대통령이 반복 사고 기업 대상 징벌적 과징금 방침을 언급한 것으로 보도됨",
        "신용정보(연소득·한도) 유출 시 보이스피싱·대출사기 악용 우려로 고객 문의·민원 증가 가능, 업권 전반 신뢰 훼손",
        "점검·대응 인력 투입, 제휴 모집인 채널 일시 제한 시 영업 영향 가능. 규모는 조사 중"],
       ["외부 노출 API 인벤토리 확인, 모집인·제휴 조회 기능 인증·권한 점검, 공유 공격 IP 차단 (정보보호부·보안관제)",
        "API 게이트웨이 단일 권한검증 표준·호출량 제한·이상 탐지 도입, 점검 결과 CEO·이사회 보고 (정보보호부·IT기획, 책임 임원 지정)",
        "API 보안 자동 테스트(DAST)·AI 활용 공격 시뮬레이션(레드팀) 정례화, 예산·인력(API 보안 전담 인력, 점검 도구) 확보 승인 필요"],
       ["[T1] 금융위원회 / 금융권 「긴급 상황대응 회의」 개최 / 2026-10-02 / https://www.fsc.go.kr/no010101/87869",
        "[T4] 한국일보 / 신한은행 해킹에 2만5000명 개인정보 유출… 금감원 긴급 현장조사 / 2026-10-01 / https://www.hankookilbo.com/news/article/A2026100115320000186",
        "[T4] 파이낸셜뉴스 / 금융사 7곳 뚫렸다… 개인 6만6천명·법인 2200건 정보 유출 / 2026-10-04 / https://fnnews.com/news/202610041822434160",
        "[T3] 데일리시큐 / ARTEX로 금융권 광범위 공격…‘API 권한검증’ 허점 / 2026-10 (일자 미표기) / https://www.dailysecu.com/news/articleView.html?idxno=208718",
        "[T4] 뉴스1 / 李대통령, 금융권 해킹 사건 \"엄중한 인식…철저 조사·대책 마련\" 지시 / 2026-10-04 / https://www.news1.kr/politics/president/6310120"])

report("Citrix NetScaler ADC·Gateway 제로데이 RCE 취약점 분석 및 대응 리포트",
       ["2026-09-27 (벤더 패치·CISA KEV 등재), 이후 2026-10-04 CVE-2026-88779 KEV 등재",
        "악용 건수 미공개(조사 중). CISA는 전 세계 악용을 확인. 당사 보유·영향 여부는 확인 중",
        "해당 없음(취약점 건). 영향: 장비 제어권, 인증 서비스(VPN·SSO) 가용성"],
       ["인증 없는 입력 검증 오류 기반 원격 명령 실행(T1190), 백도어 설치(T1505.003 보도 기반). 기술적 원인: 입력 검증 미흡 / 관리적 원인: 인터넷 노출 장비 패치·점검 주기",
        "Citrix 보안 공지(CTX697096)와 CISA 경고. 관리자들이 장비를 차단한 뒤 벤더가 제로데이를 인정(SecurityWeek)",
        "패치 이전부터 악용(패치 이전 장비 대상 악용 보도). 정확한 시점은 조사 중"],
       ["벤더 패치 번들 제공, CISA KEV 3일 기한 부여(2차 출처 기준)",
        "패치 후에도 백도어 잔존 가능성이 보도됨. 침해 흔적 점검 병행 필요",
        "CISA 경고 발령. 국내 KISA 보안공지: 해당 기간 확인되지 않음"],
       ["전자금융감독규정·개인정보 보호법상 안전조치 의무(조문 확인 필요). 침해 시 신고·통지 의무 발생 가능",
        "VPN·SSO 장애 시 임직원·상담 업무 중단 및 고객 서비스 영향 가능",
        "긴급 패치·포렌식 비용, 인증 인프라 재구축 비용 가능"],
       ["NetScaler 자산 전수 식별, CTX697096 패치, 웹셸·관리 로그 점검, 자격증명·세션 초기화 검토 (인프라보안팀)",
        "인터넷 노출 장비 패치 SLA(KEV 등재 시 72시간) 규정화 및 점검 체계 구축 (정보보호부)",
        "외부 접속 인프라의 Zero Trust(신뢰 비전제) 전환 로드맵 수립, 예산 반영 (CIO·CISO)"],
       ["[T1] CISA / Critical Zero-Day Vulnerabilities Exploited in Citrix NetScaler ADC, Gateway / 2026-09-27 / https://www.cisa.gov/news-events/alerts/2026/09/27/critical-zero-day-vulnerabilities-exploited-citrix-netscaler-adc-gateway",
        "[T3] SecurityWeek / Citrix Confirms 2 NetScaler Zero-Days After Admins Pulled the Plug / 2026-09 (일자 미표기) / https://www.securityweek.com/citrix-confirms-2-netscaler-zero-days-after-admins-pulled-the-plug/",
        "[T3] SecurityWeek / Exploitation of Citrix NetScaler Zero-Day Hits Appliances Patched Days Earlier / 2026-10 (일자 미표기) / https://www.securityweek.com/exploitation-of-citrix-netscaler-zero-day-hits-appliances-patched-days-earlier/",
        "[T2] watchTowr / Citrix NetScaler Zero-Day RCE FAQ / 2026-09 (일자 미표기) / https://watchtowr.com/intelligence/citrix-netscaler-zero-day-vulnerabilities-faq/"])

report("FortiMail 경로 순회 제로데이(CVE-2026-104286) 분석 및 대응 리포트",
       ["2026-10-01 (Fortinet 권고 FG-IR-26-175, CISA KEV 등재)",
        "악용 건수 미공개(조사 중). 당사 보유 여부는 확인 중",
        "해당 없음(취약점 건). 영향: 메일 게이트웨이 장악 및 메일 데이터"],
       ["경로 순회·NULL 바이트 처리 미흡에 의한 임의 파일 쓰기(T1190). 기술적 원인: 입력 경로 검증 미흡 / 관리적 원인: 관리·IBE 인터페이스 외부 노출",
        "Fortinet 권고 및 CISA KEV 등재",
        "조사 중"],
       ["Fortinet 우회 조치 권고(IBE 비활성화, 관리 인터페이스 인터넷 차단). 정식 패치 미제공",
        "관리 인터페이스 접근 제한, 로그·파일 변경 모니터링 필요",
        "CISA 경고 발령. 국내 KISA 보안공지: 해당 기간 확인되지 않음"],
       ["메일 경유 개인(신용)정보 유출 시 개인정보 보호법·신용정보법상 통지 의무 가능(조문 확인 필요)",
        "메일 유출 시 고객·제휴사 신뢰 저하 및 피싱 악용 가능",
        "패치 미제공 기간 중 메일 기능 제한 시 업무 영향 가능"],
       ["영향 버전 확인, IBE 비활성화, 관리 인터페이스 외부 차단 (메일보안 담당)",
        "메일 게이트웨이 노출면 축소와 SEG(보안 이메일 게이트웨이) 이중화·패치 모니터링 체계 정립 (정보보호부)",
        "메일 보안 아키텍처 재검토 및 클라우드 메일 보안 병행 도입 검토, 예산 반영 (CIO·CISO)"],
       ["[T2] Truesec / CVE-2026-104286: FortiMail Path Traversal Vulnerability Actively Exploited / 2026-10 (일자 미표기) / https://www.truesec.com/hub/blog/cve-2026-104286-fortimail-path-traversal-vulnerability-actively-exploited",
        "[T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog / 2026-10-01 / https://www.cisa.gov/news-events/alerts/2026/10/01/cisa-adds-one-known-exploited-vulnerability-catalog",
        "[T3] SecurityOnline / FortiMail Path Traversal Flaw CVE-2026-104286 Exploited in the Wild / 2026-10 (일자 미표기) / https://securityonline.info/fortimail-vulnerability-cve-2026-104286/"])

# ------------------------------------------------------------------ 기타 수록 건 출처 (SECTION 2 전용)
h2("부록 — SECTION 2 수록 건 중 상세 미기재 항목 출처")
for s in [
    "[T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog (Citrix NetScaler CVE-2026-88779) / 2026-10-04 / https://www.cisa.gov/news-events/alerts/2026/10/04/cisa-adds-one-known-exploited-vulnerability-catalog",
    "[T1] CISA / CISA Adds Two Known Exploited Vulnerabilities to Catalog (Zammad CVE-2026-102489/102490) / 2026-10-02 / https://www.cisa.gov/news-events/alerts/2026/10/02/cisa-adds-two-known-exploited-vulnerabilities-catalog",
    "[T2] Symantec / Warlock Ransomware Attackers Hit Water and Telecom Operators / 2026-10-02 / https://www.security.com/threat-intelligence/warlock-ransomware-critical-infrastructure",
    "[T3] BleepingComputer / Warlock ransomware breach SharePoint in water, telecom operator attacks / 2026-10-02 / https://www.bleepingcomputer.com/news/security/warlock-ransomware-breach-sharepoint-in-water-telecom-operator-attacks/",
    "[T3] Cybersecurity News / Critical GitLab AI Gateway Vulnerability Enables Remote Code Execution Attacks (CVE-2026-90970) / 2026-10 (일자 미표기) / https://cybersecuritynews.com/gitlab-ai-gateway-vulnerability/",
    "[T3] BleepingComputer / Police dismantle KillSec ransomware gang allegedly led by 16-year-old / 2026-10-01 / https://www.bleepingcomputer.com/news/security/police-dismantle-killsec-ransomware-gang-allegedly-led-by-16-year-old/",
    "[T3] The Hacker News / Police Arrest 16-Year-Old Suspected of Running KillSec / 2026-10 (일자 미표기) / https://thehackernews.com/2026/10/police-arrest-16-year-old-suspected-of.html",
    "[T2] AhnLab ASEC / Ransom & Dark Web Issues Week 1, October 2026 / 2026-10 (1주차) / https://asec.ahnlab.com/en/95669/",
    "[T1] 금융위원회 / 금융권 「긴급 상황대응 회의」 개최 / 2026-10-02 / https://www.fsc.go.kr/no010101/87869",
]:
    bullet(" " + s)

# ------------------------------------------------------------------ 자체검증
h1("자체검증 결과")
box([
    "① 논리적 비약 여부: 없음. 단, 당사 시스템의 영향 여부(NetScaler·FortiMail 보유, 모집인 조회 기능 존재)는 확인되지 않아 모두 조건부로 기술함.",
    "② 추측성 표현 여부: 없음. 공격 주체 귀속은 하지 않음(ARTEX는 도구 흔적이며 공격 주체 아님). 금융권 시사점 섹션의 시나리오만 조건절 형식으로 전개함.",
    "③ 실무 적용 가능성: 적합. 다만 NetScaler 영향 버전 범위, Citrix·FortiMail 벤더 원문 및 금융보안원 권고 원문은 직접 열람하지 못해 실행 전 원문 확인 필요.",
    "④ 출처 검증 상태: ✅ 확인 12건 / ⚠️ 미확인 1건 (ASEC 다크웹 동향, 단일 소스 T2).",
    "⑤ 날짜·기간 기준 검증: KST 기준 2026-09-28 ~ 2026-10-05 충족. 예외: Citrix 제로데이 공지는 2026-09-27(미국 기준 CISA 공지, 기간 직전이나 후속 악용 보도 및 CVE-2026-88779 KEV가 기간 내라 포함), "
    "Warlock 캠페인 활동은 최근 2개월이나 공개일 2026-10-02, 무신사·강남언니 개별 사고일은 기간 이전 가능성(보도일 2026-09-28 기준 포함). 일자 미표기 출처는 '2026-10 (일자 미표기)'로 표기함.",
    "※ 방법상 한계: 실행 환경의 네트워크 정책으로 대부분의 원문 사이트(CISA·금융위·언론사 등) 직접 열람이 차단되어 웹 검색 결과 요약과 복수 출처 대조에 의존함. "
    "일부 출처 URL은 검색 결과에 표시된 주소이며 본문 전체를 열람하지 않음. 일부 수치(CVSS, 패치 기한)는 2차 출처 기준이므로 벤더 원문 대조가 필요함.",
], fill="F2F2F2", title="[자체검증 결과]")

# ------------------------------------------------------------------ 푸터
para("침해사고대응팀 · 생성형 AI 자동 수집·분석 · 2026-10-05", True, 9, "595959", WD_ALIGN_PARAGRAPH.CENTER, 1)
para("출처: T1 금융위원회·금융보안원(보도 인용)·금융감독원(보도 인용)·CISA / T2 Symantec·AhnLab ASEC·Truesec·watchTowr / "
     "T3 SecurityWeek·BleepingComputer·The Hacker News·데일리시큐·보안뉴스 / T4 서울신문·한국일보·파이낸셜뉴스·머니투데이·뉴데일리·뉴스1·서울경제",
     size=8, color="7F7F7F", align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("saved", OUT)
