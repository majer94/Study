# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.8)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for n, sz in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11), ("Heading 4", 10)):
    h = doc.styles[n]
    h.font.name = FONT
    h.font.size = Pt(sz)
    h.font.bold = True
    h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para(text="", bold=False, size=None, color=None, align=None, style=None, space=3):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, bold, size, color)
    p.paragraph_format.space_after = Pt(space)
    if align:
        p.alignment = align
    return p


def add_runs(p, text, bold=False, size=None, color=None):
    # **bold** inline markup
    parts = text.split("**")
    for i, t in enumerate(parts):
        if not t:
            continue
        r = p.add_run(t)
        r.bold = bold or (i % 2 == 1)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.6 + 0.6 * level)
    return p


def table(headers, rows, widths=None, hdr_color="1F2A44", size=8.5, sev_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        add_runs(c.paragraphs[0], h, bold=True, size=size, color=RGBColor(255, 255, 255))
        shade(c, hdr_color)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], str(v), size=size)
            if sev_col is not None and i == sev_col:
                if "CRITICAL" in v or "위기" in v or "1순위" in v:
                    shade(cells[i], "F8D7DA")
                elif "HIGH" in v or "경계" in v or "2순위" in v:
                    shade(cells[i], "FFE5CC")
                elif "MEDIUM" in v or "주의" in v or "3순위" in v:
                    shade(cells[i], "FFF6CC")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def box(title, lines, color="E8EEF7"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, color)
    c.text = ""
    add_runs(c.paragraphs[0], title, bold=True, size=10.5, color=NAVY)
    for l in lines:
        p = c.add_paragraph()
        add_runs(p, l, size=9.5)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "6")
    bt.set(qn("w:space"), "1")
    bt.set(qn("w:color"), "999999")
    b.append(bt)
    pPr.append(b)


# ───────────── 헤더 ─────────────
para("주간 정보보호 동향 보고서", bold=True, size=22, color=NAVY, space=0)
para("2026-09-27 ~ 2026-10-04 (KST) · 침해사고대응팀 · 정보보호부", size=10.5, space=2)
para("보고 대상: CEO, CIO, CISO, CPO, 정보보호부장 │ 작성일: 2026-10-04 │ 작성: 정보보호부 침해사고대응팀장", size=9, space=2)
hr()

# ───────────── SECTION 1 ─────────────
doc.add_heading("SECTION 1 — EXECUTIVE SUMMARY", 1)
doc.add_heading("금주 보안 기상도", 2)
table(["등급", "사유"], [
    ["🔴 위기", "국내 금융 7개사 연쇄 해킹(신한은행 25,729명 등, 2026-10-04 기준 약 6.6만 명 유출)과 Citrix NetScaler·FortiMail·Cisco SD-WAN 등 CISA KEV(Known Exploited Vulnerabilities, 알려진 악용 취약점) 등재 외부노출 장비 취약점 3종, Bitget 3.875억 달러 탈취."],
    ["🟠 경계", "금융위 긴급 상황대응 회의(2026-10-02)에서 외부노출 IT자산 전면 점검·보고 지시(카드사 3개사 참석). Zammad(고객지원 플랫폼) 취약점 2건 KEV 등재."],
    ["🟡 주의", "Warlock 랜섬웨어의 SharePoint 취약점 악용, KillSec 랜섬웨어 조직 검거, 미 국방부 DMDC 인사정보 유출, Frontline Education 제3자 SW 취약점 유출."],
], widths=[2.5, 14.5], sev_col=0)

doc.add_heading("핵심 위협 요약", 2)
para("글로벌 환경에서는 Citrix NetScaler(2026-09-27), Cisco Catalyst SD-WAN Manager(2026-09-30), Fortinet FortiMail(2026-10-01) 등 인터넷에 노출된 네트워크·메일 보안 장비의 인증 불필요(Unauthenticated) 취약점이 연이어 KEV에 등재되었고, Bitget은 제3자 보안 장비 제로데이(Zero-day, 패치 이전 취약점)로 약 3.875억 달러를 탈취당했다. 공통 패턴은 \"외부 접점 장비 → 웹쉘/백도어 → 내부망 이동\"이다. "
     "국내에서는 2026-10-01 신한은행 고객 25,729명 정보 유출을 시작으로 KB국민은행(119명), 하나은행(89명), BNK부산은행(외주 개발직원 11명), 예가람저축은행(약 4만 명 추정), 현대캐피탈(주택대출 모집인 146명 중 일부) 등 금융 7개사에서 정보 유출이 확인되었다. 공격은 대출모집인 조회 서비스, 직원용 모바일 업무지원시스템 등 외부 접근 가능한 주변 업무 시스템에 집중되었으며, 금융당국은 동일 공격자 IP가 복수 사에서 발견되었고 AI(Artificial Intelligence, 인공지능) 활용 자동화 공격 정황이 있다고 밝혔다. "
     "규제 측면에서 금융위원회는 2026-10-02 긴급 상황대응 회의에서 외부노출 IT자산 파악, 보안취약점·인증·접근통제·침해탐지체계 전면 재점검 및 점검 후 보고를 지시했고, 2026-10-04 전 금융권을 소집했다. "
     "당사 관점의 가장 긴급한 조치 사항은 대출모집인·영업점·제휴사·직원용 모바일 등 외부 접근 가능 업무 시스템과 Citrix·FortiMail·Cisco 장비의 전수 식별 및 인증·인가 점검, 패치·완화조치 적용이다.")

doc.add_heading("⚡ 긴급조치 Top 3", 2)
for i, t in enumerate([
    "외부노출 IT자산(대출모집인·가맹점·제휴사·직원용 모바일·콜센터 연계 포함) 전수 식별 및 인증·인가·조회 한도 점검 — 정보보호부·IT기획부·채널 운영 부서",
    "Citrix NetScaler(CVE-2026-88771/88772)·Cisco SD-WAN Manager(CVE-2026-76504) 패치 적용 및 침해 흔적 점검, FortiMail(CVE-2026-104286) IBE(Identity-Based Encryption) 기능 비활성화 — 인프라팀·보안관제",
    "금융위 점검·보고 지시 대응 TF 가동 및 동일 공격자 IP·대량 조회 패턴 탐지룰 적용(금융권 공유 위협정보 반영) — 침해사고대응팀·보안관제센터",
], 1):
    para(f"{i}. {t}", space=2)
hr()

# ───────────── SECTION 2 ─────────────
doc.add_heading("SECTION 2 — 사건 및 위협 현황 리스트", 1)
para("수집 기간: 2026-09-27 ~ 2026-10-04 (KST). 총 11건. 기간 외 확인 항목(Cisco Secure Email Gateway CVE-2026-76461: 2026-09-14 공개, Oracle PeopleSoft ShinyHunters 캠페인: 2026-09-25 공개)은 수집 대상에서 제외하되 당사 자산 점검 시 참고한다.", size=8.5)
S2 = [
    ["🔴 CRITICAL", "개인정보침해", "국내 금융 7개사 연쇄 해킹·정보유출(신한·KB국민·하나·BNK·예가람·웰컴·현대캐피탈)", "국내 은행·저축은행·캐피탈", "2026-10-01~10-04", "✅", "높음"],
    ["🔴 CRITICAL", "취약점 악용", "Citrix NetScaler ADC/Gateway 제로데이 CVE-2026-88771·88772 (KEV)", "전 세계 NetScaler 이용 기관", "2026-09-27", "✅", "높음"],
    ["🔴 CRITICAL", "취약점 악용", "Fortinet FortiMail CVE-2026-104286 (KEV, 패치 미제공 보도)", "FortiMail 이용 기관", "2026-10-01", "✅", "높음"],
    ["🔴 CRITICAL", "취약점 악용", "Cisco Catalyst SD-WAN Manager 인증우회 CVE-2026-76504 (KEV)", "SD-WAN Manager 이용 기관", "2026-09-30", "✅", "높음"],
    ["🔴 CRITICAL", "APT·해킹", "Bitget 거래소 3.875억 달러 탈취(제3자 보안장비 제로데이)", "Bitget(가상자산 거래소)", "2026-10-02", "✅", "중간"],
    ["🟠 HIGH", "정책·제도", "금융위 긴급 상황대응 회의·전 금융권 소집(외부노출 자산 점검·보고 지시)", "금융권 전체(카드사 3개사 참석)", "2026-10-02", "✅", "높음"],
    ["🟠 HIGH", "취약점 악용", "Zammad 세션 고정·권한 관리 취약점 CVE-2026-102489·102490 (KEV)", "Zammad 이용 기관", "2026-10-02", "⚠️", "중간"],
    ["🟡 MEDIUM", "랜섬웨어", "Warlock(Storm-2603) SharePoint 취약점 악용 랜섬웨어 확산", "포르투갈어·스페인어권 사회기반시설·정부·교육", "2026-10-02", "✅", "낮음"],
    ["🟡 MEDIUM", "랜섬웨어", "KillSec 랜섬웨어 조직 3명 검거·유출사이트 압수(110TB)", "KillSec 피해 기관(약 1,000건 공격 연관)", "2026-09-30", "✅", "낮음"],
    ["🟡 MEDIUM", "개인정보침해", "미 국방부 DMDC 인사정보 유출(276만 명 생존자 등)", "미 국방부", "2026-10-01", "✅", "낮음"],
    ["🟡 MEDIUM", "공급망", "Frontline Education 제3자 SW 취약점 악용 직원정보 유출", "미국 K-12 학군", "2026-10-01", "⚠️", "낮음"],
]
table(["심각도", "위협유형", "사건명", "피해 주체", "발생·공개일", "교차검증", "금융권 관련성"], S2,
      widths=[2.0, 1.8, 5.2, 3.2, 2.2, 1.2, 1.4], size=8, sev_col=0)
para("Tier별 수집 현황: Tier 1 — CISA(KEV 등재 알림), 금융위원회 긴급회의 발표(보도 인용). KISA·금융보안원·개인정보보호위원회·NIST·ENISA 등은 검색 범위 내 해당 기간 신규 자료 미확인(해당 기간 특이사항 없음). Tier 2 — Mandiant, Symantec, Rapid7, Qualys. Tier 3 — The Hacker News, BleepingComputer, 보안뉴스, SecurityAffairs. Tier 4 — 한국일보, 파이낸셜뉴스, 뉴스핌, 서울신문, CNBC 등.", size=8.5)
hr()

# ───────────── SECTION 3 ─────────────
doc.add_heading("SECTION 3 — 위협 유형별 통계", 1)
table(["위협유형", "건수", "심각도 분포"], [
    ["취약점 악용", "4", "🔴 3 / 🟠 1"],
    ["랜섬웨어", "2", "🟡 2"],
    ["APT·해킹", "1", "🔴 1"],
    ["개인정보침해", "2", "🔴 1 / 🟡 1"],
    ["공급망", "1", "🟡 1"],
    ["정책·제도", "1", "🟠 1"],
    ["**총계**", "**11**", "🔴 5 / 🟠 2 / 🟡 4"],
], widths=[5, 2, 10], size=9)
hr()

# ───────────── SECTION 4 ─────────────
doc.add_heading("SECTION 4 — 주요 사건 상세 보고", 1)
para("금융권 영향도·심각도 기준 6건 선정(심각도 내림차순).", size=8.5)


def incident(title, overview, meta_req, meta_cond, sources, cross, risk, impact, action):
    doc.add_heading(title, 3)
    doc.add_heading("사건 개요", 4)
    for l in overview:
        bullet(l)
    doc.add_heading("메타데이터", 4)
    for l in meta_req + meta_cond:
        bullet(l)
    doc.add_heading("출처 및 교차확인", 4)
    for l in sources:
        bullet(l)
    bullet(cross)
    doc.add_heading("💳 금융권 시사점", 4)
    bullet("**[리스크]** " + risk)
    bullet("**[영향]** " + impact)
    bullet("**[조치]** " + action)
    hr()


incident(
    "🔴 [개인정보침해] 국내 금융 7개사 — 외부 접근 가능 주변 업무시스템 대상 AI 활용 자동화 공격 정황",
    [
        "누가: 피해 주체는 신한은행, KB국민은행, 하나은행, BNK부산은행, 예가람저축은행, 웰컴저축은행, 현대캐피탈. 공격 주체는 미공개(금융당국은 동일 공격자 IP 다수 확인).",
        "언제: 2026-10-01 신한은행 유출 공개 이후 10-02~10-04 순차 확인. 예가람저축은행은 2026-09-30 서버 접근 정황 인지.",
        "어디서: 대출모집인 조회 서비스, 직원용 모바일 업무지원시스템, 영업지원시스템 등 외부 접근 가능 주변 시스템.",
        "무엇을: 고객명·전화번호·연소득·대출 산출한도·신청상품·금리, 일부 주민등록번호(66건)·연계정보(CI, 97건) 등. 우리은행·NH농협은행은 공격을 차단했고 유출은 확인되지 않음.",
        "어떻게: LLM(Large Language Model, 대규모 언어모델)·멀티에이전트 기반 자동화 침투 도구로 인증·인가 허점을 대량 조회한 정황이 보도됨(중국어 오픈소스 도구 흔적 보도). 시스템 장악이 아닌 취약 시스템 접근 후 조회·반출 방식.",
        "얼마나: 신한 25,729명, 예가람 약 4만 명(추정), KB국민 119명(고객 99·임직원 20), 하나 89명, BNK 11명, 현대캐피탈 모집인 146명 중 일부. 개인 약 6.6만 명·법인 2,200건(파이낸셜뉴스 집계). 웰컴저축은행은 조사 중.",
    ],
    ["발생·공개일: 2026-10-01 ~ 2026-10-04",
     "공격 유형: T1190 Exploit Public-Facing Application(외부 노출 애플리케이션 악용), T1595.002 Vulnerability Scanning(취약점 스캐닝), T1119 Automated Collection(자동 수집)",
     "피해 규모: 약 6.6만 명 + 법인 2,200건(집계 시점 2026-10-04, 변동 가능)"],
    ["유출 항목: 고객명, 전화번호, 연소득, 대출 산출한도, 신청상품·금리, 주민등록번호 66건, CI 97건, 생년월일(예가람)",
     "관련 규제: 개인정보 보호법, 신용정보의 이용 및 보호에 관한 법률, 전자금융거래법 (조문 확인 필요)"],
    ["출처 1: [T4] 한국일보 / 신한은행 이어 KB국민은행도 119명 개인정보 유출… 당국, 은행권 긴급 소집 / 2026-10-02 / https://www.hankookilbo.com/news/article/A2026100214560000654",
     "출처 2: [T4] 뉴스핌 / 금융당국, 신한·국민 잇단 해킹에 긴급회의…\"외부노출 전산시스템 전면 점검\" / 2026-10-02 / https://www.newspim.com/news/view/20261002001168",
     "출처 3: [T3] 보안뉴스 / [은행이 뚫린다] \"AI 혁신 전 기본기부터\"... 초보적 공격에 신한은행 고객 정보 유출 / 2026-10 / https://www.boannews.com/news/articleView.html?idxno=146174",
     "출처 4: [T4] 이투데이 / 이억원 \"금융권 해킹 사고, AI 활용 추정… 동일 공격자 IP 여러 곳 발견\" / 2026-10-04 / https://www.etoday.co.kr/news/view/2632359"],
    "교차확인: ✅ 확인 (4개 소스 / 최고 신뢰도: T3 보안뉴스, 금융당국 발표 인용)",
    "확인된 사실: 금융당국이 7개사 유출과 동일 공격자 IP를 확인했고, 직접 사기거래에 쓸 수 있는 민감정보 유출은 확인되지 않았다고 밝혔다. 당사가 동일 유형 사고를 겪는 경우 개인정보 보호법(유출 통지·신고), 신용정보법, 전자금융거래법상 안전성 확보 의무 위반 여부 조사와 감독당국 검사·제재 대상이 될 수 있다(조문·제재 수준은 확인 필요). 금융위가 점검 후 보고를 지시해 보고 미비 자체도 감독 리스크다.",
    "대출모집인 조회·직원용 모바일 등 외부 접근 시스템에 인증·인가 허점이 있었다는 사실이 확인된 만큼, 당사 모집인·제휴사·영업지원 채널에 유사 허점이 있는 경우 소액 조회를 반복하는 자동화 공격으로 고객 신용정보가 대량 반출되고 민원·감독 대응 비용이 발생할 수 있다.",
    "(1) 외부 접근 가능 업무시스템 전수 식별 후 인증·인가·조회 횟수 제한·마스킹 점검 — 정보보호부·채널 운영 부서. (2) 해외 IP·비정상 조회량·봇 패턴 탐지룰 적용 및 공유 IP 차단 — 보안관제센터.",
)

incident(
    "🔴 [취약점 악용] Citrix NetScaler ADC·Gateway — 인증 불필요 RCE 제로데이 2종, 웹쉘 설치·내부망 이동",
    [
        "누가: 공격 주체 미공개. 피해 주체는 NetScaler ADC·Gateway 이용 기관.",
        "언제: 2026-09-27 Citrix 취약점 8건 공개, CISA가 같은 날 KEV 등재. 공개 이전부터 제로데이로 악용(수 주간).",
        "어디서: 전 세계. Shadowserver 집계 기준 23,000대 이상이 인터넷 노출.",
        "무엇을: NetScaler ADC·Gateway(VPN 가상서버 포함).",
        "어떻게: CVE-2026-88771은 입력 검증 결함으로 기본 설정에서 명령 실행, CVE-2026-88772는 기본 활성화된 DTLS(Datagram TLS) 경로의 버퍼 경계 위반. 악용 후 맞춤형 웹쉘·터널링 악성코드 설치, root 권한 획득, 자격증명 탈취 및 내부망 확산이 보고됨.",
        "얼마나: 구체 피해 기관 수 조사 중. 연방기관은 2026-09-30까지 패치 또는 장비 분리 명령(BOD 26-04).",
    ],
    ["발생·공개일: 2026-09-27",
     "공격 유형: T1190 Exploit Public-Facing Application, T1505.003 Web Shell(웹쉘), T1572 Protocol Tunneling(프로토콜 터널링)",
     "피해 규모: 조사 중(노출 장비 23,000대 이상)"],
    ["CVE: CVE-2026-88771, CVE-2026-88772 (CVSS v4 9.5 / Critical)",
     "영향 버전: NetScaler ADC·Gateway(상세 버전은 벤더 권고 확인 필요)",
     "CISA KEV 등재 여부: 예(2026-09-27)",
     "패치 기한: 2026-09-30 (미 연방기관 기준)"],
    ["출처 1: [T1] CISA / Critical Zero-Day Vulnerabilities Exploited in Citrix NetScaler ADC, Gateway / 2026-09-27 / https://www.cisa.gov/news-events/alerts/2026/09/27/critical-zero-day-vulnerabilities-exploited-citrix-netscaler-adc-gateway",
     "출처 2: [T2] Rapid7 / Zero-Day Exploitation of Citrix NetScaler ADC and Gateway: CVE-2026-88771 and CVE-2026-88772 / 2026-09 / https://www.rapid7.com/blog/post/etr-zero-day-exploitation-of-citrix-netscaler-adc-and-gateway-cve-2026-88771-and-cve-2026-88772/",
     "출처 3: [T3] Help Net Security / Citrix NetScaler RCE zero-days exploited globally for weeks / 2026-09-28 / https://www.helpnetsecurity.com/2026/09/28/citrix-netscaler-rce-zero-days-exploited-for-weeks-cve-2026-88771-cve-2026-88772/"],
    "교차확인: ✅ 확인 (3개 소스 / 최고 신뢰도: T1 CISA)",
    "KEV 등재와 수 주간의 제로데이 악용이 확인되었으므로, 당사가 NetScaler를 VPN·가상화 접속에 사용하는 경우 전자금융거래법상 접근통제·침해사고 보고 의무와 개인정보 보호법상 안전조치 의무(조문 확인 필요)가 문제될 수 있다.",
    "공개 이전부터 악용된 사실이 확인되었으므로, 당사 NetScaler가 패치 전 노출되어 있었던 경우 웹쉘·자격증명 탈취 후 내부망 이동으로 승인·고객정보 시스템까지 영향이 확산될 수 있다.",
    "(1) 벤더 패치 즉시 적용 및 패치 전후 웹쉘·비정상 프로세스 점검, 관리자·VPN 계정 재설정 — 인프라팀. (2) 패치 불가 장비는 외부 노출 차단 — 네트워크보안팀.",
)

incident(
    "🔴 [취약점 악용] Fortinet FortiMail — IBE 기능 경로조작으로 인증 없는 임의 파일 쓰기(패치 미제공 보도)",
    [
        "누가: 공격 주체 미공개. 피해 주체는 FortiMail 이용 기관.",
        "언제: 2026-10-01 CISA KEV 등재.",
        "어디서: FortiMail 웹 인터페이스의 IBE(Identity-Based Encryption, 신원기반 암호화) 기능.",
        "무엇을: FortiMail 8.0.0~8.0.1, 7.6.0~7.6.6, 7.4.0~7.4.8, 7.2.0~7.2.9.",
        "어떻게: 경로 조작(CWE-22)과 NULL 바이트 처리 미흡(CWE-158)을 결합해 인증 없이 조작된 HTTP(S) 요청으로 임의 파일 기록.",
        "얼마나: 피해 규모 조사 중. 보도 기준 공식 패치 미제공이며 IBE 비활성화가 임시 조치.",
    ],
    ["발생·공개일: 2026-10-01",
     "공격 유형: T1190 Exploit Public-Facing Application",
     "피해 규모: 조사 중"],
    ["CVE: CVE-2026-104286 (CVSS 9.8 / Critical)",
     "영향 버전: FortiMail 7.2.0~7.2.9, 7.4.0~7.4.8, 7.6.0~7.6.6, 8.0.0~8.0.1",
     "CISA KEV 등재 여부: 예(2026-10-01)",
     "패치 기한: 확인 중(CISA 지정 기한 미확인; 벤더 패치 미제공 보도)"],
    ["출처 1: [T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog / 2026-10-01 / https://www.cisa.gov/news-events/alerts/2026/10/01/cisa-adds-one-known-exploited-vulnerability-catalog",
     "출처 2: [T2] runZero / Fortinet FortiMail CVE-2026-104286: Find impacted appliances / 2026-10 / https://www.runzero.com/blog/fortinet-fortimail/",
     "출처 3: [T3] Rescana / Fortinet FortiMail CVE-2026-104286 Actively Exploited / 2026-10 / https://www.rescana.com/post/fortinet-fortimail-cve-2026-104286-actively-exploited-critical-path-traversal-and-null-byte-vulnerability-alert"],
    "교차확인: ✅ 확인 (3개 소스 / 최고 신뢰도: T1 CISA)",
    "패치 없이 악용이 확인된 사실에 비추어, 당사가 FortiMail을 메일 보안에 사용하는 경우 조치 지연 시 전자금융거래법·정보보호 관련 감독 점검에서 취약점 관리 미흡으로 지적될 수 있다(조문 확인 필요).",
    "인증 없는 파일 쓰기가 가능한 장비가 이메일 경로에 있는 경우 웹쉘 설치를 통한 메일·고객 안내 발송 시스템 장악과 피싱 발송 악용이 발생할 수 있다.",
    "(1) 해당 버전 보유 여부 확인 후 IBE 기능 비활성화 및 관리 인터페이스 외부 차단 — 메일보안 담당. (2) 벤더 패치 공지 모니터링·적용 및 웹쉘·비정상 파일 점검 — 인프라팀.",
)

incident(
    "🔴 [취약점 악용] Cisco Catalyst SD-WAN Manager — URI 인코딩 처리 결함으로 인증 없이 관리자 권한 획득",
    [
        "누가: 공격 주체 미공개. 피해 주체는 SD-WAN Manager 이용 기관.",
        "언제: 2026-09-30 Cisco 보안공지 및 CISA KEV 등재(Cisco는 2026-09 중 악용 인지).",
        "어디서: Catalyst SD-WAN Manager(전 배포 형태 영향).",
        "무엇을: SD-WAN 중앙 관리 시스템의 관리자 권한.",
        "어떻게: HTTP 요청 URI 인코딩 처리 결함(hex encoding)을 악용한 인증 우회.",
        "얼마나: 피해 규모 조사 중. 패치 배포됨.",
    ],
    ["발생·공개일: 2026-09-30",
     "공격 유형: T1190 Exploit Public-Facing Application, T1078 Valid Accounts(관리자 권한 획득)",
     "피해 규모: 조사 중"],
    ["CVE: CVE-2026-76504 (CVSS 9.8 / Critical)",
     "영향 버전: Catalyst SD-WAN Manager 전 배포(고정 릴리스는 Cisco 공지 확인 필요)",
     "CISA KEV 등재 여부: 예(2026-09-30)",
     "패치 기한: 확인 중(CISA 지정 기한 미확인)"],
    ["출처 1: [T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog / 2026-09-30 / https://www.cisa.gov/news-events/alerts/2026/09/30/cisa-adds-one-known-exploited-vulnerability-catalog",
     "출처 2: [T2] Qualys / Cisco Catalyst SD-WAN Manager Authentication Bypass Vulnerability Exploited in Attacks (CVE-2026-76504) / 2026-09-30 / https://threatprotect.qualys.com/2026/09/30/cisco-catalyst-sd-wan-manager-authentication-bypass-vulnerability-exploited-in-attacks-cve-2026-76504/",
     "출처 3: [T3] The Hacker News / CISA Adds Exploited Cisco Catalyst SD-WAN Auth Bypass to KEV / 2026-10 / https://thehackernews.com/2026/10/cisa-adds-exploited-cisco-catalyst-sd.html"],
    "교차확인: ✅ 확인 (3개 소스 / 최고 신뢰도: T1 CISA)",
    "KEV 등재가 확인되었으므로, 당사가 SD-WAN으로 영업점·IDC·클라우드를 연결하는 경우 관리 평면 장악 시 전자금융 서비스 연결 구간 전체가 영향받아 전자금융거래법상 안정성 확보 의무 이행 점검 대상이 될 수 있다(조문 확인 필요).",
    "관리자 권한이 탈취되는 경우 SD-WAN 정책 변경을 통해 트래픽 우회·망분리 경계 약화가 발생할 수 있다.",
    "(1) Cisco 패치 적용 및 Manager 관리 인터페이스 외부 노출 차단·관리자 계정 점검 — 네트워크팀. (2) 정책 변경 이력 감사 및 알림 설정 — 보안관제센터.",
)

incident(
    "🔴 [APT·해킹] Bitget — 제3자 보안장비 제로데이로 핫·웜 월렛 3.875억 달러 탈취",
    [
        "누가: 피해 주체는 Bitget. 공격 주체는 미확인(Bitget CEO는 북한 해커로 주장했으나 독립 검증 미확인, 본 보고서는 귀속하지 않음).",
        "언제: 최초 접근 2026-08-31 이상, 탈취 실행 2026-09-25 자정 이후, 원인 공개 2026-10-01~10-02.",
        "어디서: Bitget 핫·웜 월렛 및 월렛 작업 서버.",
        "무엇을: 이더리움·XRP·Arbitrum·Avalanche·Optimism·BSC·Base 등 자산 3.875억 달러.",
        "어떻게: 제3자 보안장비 2종의 제로데이로 침투, 웹쉘 설치·C2 구성 후 월렛 서버로 이동, 악성 패키지와 거래 데이터를 위조하는 맞춤 출금 도구로 승인을 유발.",
        "얼마나: 3.875억 달러. CEO는 \"상당액 회수를 기대하지 않는다\"고 CNBC에 밝힘.",
    ],
    ["발생·공개일: 2026-10-02 (탈취: 2026-09-25)",
     "공격 유형: T1190 Exploit Public-Facing Application, T1505.003 Web Shell, T1021 Remote Services(내부 이동), T1565 Data Manipulation(거래 데이터 조작)",
     "피해 규모: 3.875억 달러"],
    ["CVE: 미공개(제3자 보안장비 제로데이 2종, 제품명 조사 중)",
     "관련 규제: 전자금융거래법, 특정 금융거래정보의 보고 및 이용 등에 관한 법률(가상자산 관련, 조문 확인 필요)"],
    ["출처 1: [T3] The Hacker News / Bitget hacked via zero-day in third-party security products / 2026-10 / https://x.com/TheHackersNews/status/2105528856048496964",
     "출처 2: [T4] CNBC / Bitget 'not expecting to recover a lot' from $388 million hack, CEO tells CNBC / 2026-10-02 / https://www.cnbc.com/2026/10/02/bitget-crypto-stolen-hack-recovery.html",
     "출처 3: [T2] Mandiant(조사 주체, 보도 인용) / daily.dev 요약 / 2026-10 / https://daily.dev/posts/bitget-hacked-via-zero-day-in-third-party-security-products-tiafv2oyn"],
    "교차확인: ✅ 확인 (2개 독립 소스: The Hacker News, CNBC / 최고 신뢰도: T3)",
    "직접 금전 피해가 확인되었고, 보안장비 자체가 침투 경로였다는 사실이 확인되었다. 당사가 동일 유형 장비(VPN·WAF·메일보안)를 운영하는 경우 제3자 보안제품 관리와 이상거래 통제가 전자금융거래법상 안전성 의무 점검 대상이 될 수 있다(조문 확인 필요).",
    "보안장비 제로데이로 내부망에 진입한 공격자가 출금·승인 서버에 접근하는 경우 거래 데이터 위조로 정상 승인 절차가 우회되어 직접 금전 피해로 이어질 수 있다.",
    "(1) 승인·정산·결제 서버 접근경로의 보안장비 의존도 점검 및 관리망 분리 — 인프라보안팀. (2) 고액·비정상 승인 요청에 대한 독립 검증(이중 승인) 적용 점검 — 승인·정산 시스템 담당.",
)

incident(
    "🟠 [정책·제도] 금융위원회 — 긴급 상황대응 회의 및 전 금융권 소집, 외부노출 IT자산 점검·보고 지시",
    [
        "누가: 금융위원회(신진창 사무처장 주재), 금융감독원, 금융보안원, 주요 은행 6개사, 카드사 3개사, 은행연합회, 여신금융협회.",
        "언제: 2026-10-02 긴급 상황대응 회의, 2026-10-04 금융위원장 발언 및 전 금융권 소집.",
        "어디서: 정부서울청사 등 금융권 전체.",
        "무엇을: 외부노출 IT자산·서비스 파악, 보안취약점·인증·접근통제·침해활동 탐지체계 재점검, 위협정보 신속 공유 및 공동대응.",
        "어떻게: 점검 후 보고 지시(보도: \"은행·카드 보안 점검 후 보고\").",
        "얼마나: 금융권 전체 대상. 금융위원장은 유사 사고 시 엄정 대응 방침을 표명.",
    ],
    ["발생·공개일: 2026-10-02 ~ 2026-10-04",
     "공격 유형: 해당 없음(정책·제도)",
     "피해 규모: 해당 없음"],
    ["관련 규제: 전자금융거래법, 전자금융감독규정(조문 확인 필요)"],
    ["출처 1: [T4] 뉴스핌 / 금융당국, 신한·국민 잇단 해킹에 긴급회의…\"외부노출 전산시스템 전면 점검\" / 2026-10-02 / https://www.newspim.com/news/view/20261002001168",
     "출처 2: [T4] 서울신문 / 해킹 사고 잇따르자…당국 \"외부 노출 전산시스템 빠짐없이 점검하라\" / 2026-10-02 / https://www.seoul.co.kr/news/economy/finance/2026/10/02/20261002500245",
     "출처 3: [T4] 비즈워치 / 신한·KB 정보유출에 금융당국 긴급회의…\"은행·카드 보안 점검후 보고\" / 2026-10-02 / https://news.bizwatch.co.kr/article/finance/2026/10/02/0044",
     "출처 4: [T4] 이투데이 / 이억원 \"금융권 해킹 사고, AI 활용 추정…\" / 2026-10-04 / https://www.etoday.co.kr/news/view/2632359"],
    "교차확인: ✅ 확인 (4개 소스 / 최고 신뢰도: T4, 금융위 발표 인용)",
    "점검·보고 지시가 확인되었으므로 당사의 점검 누락이나 보고 지연은 감독 검사에서 관리 소홀로 평가될 수 있으며, 유사 사고 시 엄정 대응 방침에 따라 제재 수위가 가중될 수 있다(조문·제재 수준 확인 필요).",
    "금융위가 외부노출 자산 전수 점검을 요구한 사실에 비추어, 당사 자산 인벤토리가 불완전한 경우 점검 보고 품질 문제와 미식별 자산 침해 위험이 동시에 발생할 수 있다.",
    "(1) 점검·보고 대응 TF 구성 및 감독당국 요청 서식·기한 확인 — 정보보호부·감사/준법. (2) 외부 공개 자산 스캔 결과를 인벤토리와 대조 — 정보보호부·IT인프라.",
)

# ───────────── SECTION 5 ─────────────
doc.add_heading("SECTION 5 — 즉시 조치 권고", 1)
table(["우선순위", "조치 항목", "담당 부서/시스템", "완료 기한", "근거 사건"], [
    ["🔴 1순위", "NetScaler(CVE-2026-88771/88772) 패치 적용 및 웹쉘·관리자 계정 침해 점검", "인프라팀·VPN/ADC", "**즉시**", "Citrix NetScaler"],
    ["🔴 1순위", "Cisco SD-WAN Manager 패치, FortiMail IBE 비활성화 및 관리 인터페이스 외부 차단", "네트워크팀·메일보안", "**즉시**", "Cisco SD-WAN / FortiMail"],
    ["🔴 1순위", "대출모집인·제휴사·직원용 모바일 등 외부 접근 시스템 인증·인가·조회 한도 긴급 점검, 비정상 조회 탐지룰 적용", "정보보호부·채널 운영·보안관제", "**즉시**", "국내 금융권 연쇄 해킹"],
    ["🟠 2순위", "외부 노출 IT자산 전수 식별 및 금융위 점검·보고 대응(인벤토리 대조)", "정보보호부·IT인프라·준법", "**금주 내**", "금융위 긴급회의"],
    ["🟠 2순위", "보안장비·제3자 제품(VPN·WAF·메일)의 관리망 분리 및 승인·정산 서버 접근경로 점검", "인프라보안팀", "**금주 내**", "Bitget"],
    ["🟡 3순위", "SharePoint 등 외부 노출 협업 SW 패치 현황 점검, 제3자 SW 취약점 통지 체계 정비", "IT인프라·협력사관리", "**이달 내**", "Warlock / Frontline"],
], widths=[2.0, 7.0, 3.8, 1.8, 3.0], size=8.5, sev_col=0)
hr()

# ───────────── SECTION 6 ─────────────
doc.add_heading("SECTION 6 — 사고 개별 분석 리포트", 1)


def report(title, a, b, c, d1, d2, _x, _y, _z, srcs):
    doc.add_heading(title, 2)
    doc.add_heading("1. 사고 개요 및 규모", 3)
    for l in a:
        bullet(l)
    doc.add_heading("2. 사고 원인 및 경위", 3)
    for l in b:
        bullet(l)
    doc.add_heading("3. 현재 대응 상황", 3)
    for l in c:
        bullet(l)
    doc.add_heading("4. 경영진 제언", 3)
    para("**비즈니스 리스크**", space=1)
    for l in d1:
        bullet(l)
    para("**대응 전략**", space=1)
    for l in d2:
        bullet(l)
    doc.add_heading("5. 관련 출처", 3)
    for l in srcs:
        bullet(l)
    hr()


report(
    "[보고] 국내 금융권 연쇄 해킹·정보유출 사고 분석 및 대응 리포트",
    ["발견 시점: 2026-10-01(신한은행 공개, 금융당국·언론 보도), 예가람저축은행은 2026-09-30 접근 정황 인지",
     "유출·피해 규모: 개인 약 6.6만 명·법인 2,200건(2026-10-04 집계, 변동 가능). 신한 25,729명 등",
     "피해 항목: 고객명, 전화번호, 연소득, 대출 산출한도, 신청상품·금리, 일부 주민등록번호·CI"],
    ["공격 유형: T1190 Exploit Public-Facing Application, T1595.002 Vulnerability Scanning, T1119 Automated Collection(자동화 도구 활용 정황은 보도 기준, 공격 주체 미공개)",
     "발견 배경: 금융사별 자체 인지 및 금융당국 취합. 당국이 동일 공격자 IP 다수 확인",
     "발생 추정 시점: 조사 중(예가람저축은행 2026-09-30 인지)"],
    ["즉시 대응: 금융위 긴급 상황대응 회의(2026-10-02), 전 금융권 소집(2026-10-04). 우리은행·NH농협은행은 공격 차단 보도",
     "보안 강화: 외부노출 자산 점검, 인증·접근통제 강화 지시. 개별사 조치는 조사 중",
     "유관기관 신고: 금융당국 주도 대응 확인. 개인정보보호위원회·KISA 신고 세부는 미공개"],
    ["법적·규제: 개인정보 보호법, 신용정보법, 전자금융거래법상 안전조치·통지·신고 의무 및 감독 검사 대상(조문·제재 수준 확인 필요)",
     "평판·고객: 대출·신용정보 유출에 따른 민원·스미싱 우려, 다음 피해 금융사로 거론되는 데 따른 평판 부담",
     "재무·운영: 점검·포렌식 비용, 모집인·제휴 채널 운영 제한 가능성"],
    ["즉시(24시간 이내): 모집인·직원용 모바일·영업지원 외부 접근 경로 인증·인가 점검 및 비정상 조회 차단 — 채널 시스템·보안관제",
     "단기(1~2주 이내): 외부노출 자산 인벤토리 정비, 조회 한도·마스킹·접근 로깅 표준화 — 정보보호부·IT기획부",
     "중장기(1분기 이내): 외부 접점 시스템 Zero Trust(제로트러스트) 적용, 자동화 공격 대응용 봇·이상행위 탐지 체계 구축 — 예산·인력 증원 검토 필요"],
    None, None, None,
    ["[T4] 한국일보 / 신한은행 이어 KB국민은행도 119명 개인정보 유출… 당국, 은행권 긴급 소집 / 2026-10-02 / https://www.hankookilbo.com/news/article/A2026100214560000654",
     "[T4] 뉴스핌 / 금융당국, 신한·국민 잇단 해킹에 긴급회의 / 2026-10-02 / https://www.newspim.com/news/view/20261002001168",
     "[T4] 이투데이 / 이억원 \"금융권 해킹 사고, AI 활용 추정\" / 2026-10-04 / https://www.etoday.co.kr/news/view/2632359",
     "[T3] 보안뉴스 / [은행이 뚫린다] 초보적 공격에 신한은행 고객 정보 유출 / 2026-10 / https://www.boannews.com/news/articleView.html?idxno=146174",
     "[T3] 보안뉴스 / [은행이 뚫린다] 현대캐피탈 해킹 공격... 주택대출 모집인 146명 정보 유출 / 2026-10 / https://www.boannews.com/news/articleView.html?idxno=146187"],
)

report(
    "[보고] 외부노출 네트워크·메일 장비 KEV 취약점(Citrix·FortiMail·Cisco SD-WAN) 분석 및 대응 리포트",
    ["발견 시점: Citrix 2026-09-27, Cisco SD-WAN 2026-09-30, FortiMail 2026-10-01(각 CISA KEV 등재)",
     "유출·피해 규모: 조사 중(NetScaler 노출 장비 23,000대 이상)",
     "피해 항목: 미공개(웹쉘·자격증명 탈취 보고)"],
    ["공격 유형: T1190, T1505.003, T1572, T1078",
     "발견 배경: 벤더 보안공지와 CISA KEV 등재",
     "발생 추정 시점: NetScaler는 공개 수 주 전부터 악용(Help Net Security 보도)"],
    ["즉시 대응: 벤더 패치(Citrix·Cisco), FortiMail은 공식 패치 미제공 보도로 IBE 비활성화 안내",
     "보안 강화: 웹쉘·비정상 프로세스 점검, 관리 인터페이스 외부 차단 권고",
     "유관기관 신고: CISA 긴급 지침(BOD 26-04, 미 연방기관 대상). 국내 신고 해당 여부 조사 중"],
    ["법적·규제: 전자금융거래법상 접근통제·취약점 관리 의무(조문 확인 필요), 침해 시 개인정보 보호법 유출 통지",
     "평판·고객: 메일·VPN 장악 시 고객 대상 피싱 발송 및 서비스 신뢰 훼손",
     "재무·운영: 긴급 패치 중 서비스 중단, 장비 교체·포렌식 비용"],
    ["즉시(24시간 이내): 3종 장비 보유 여부·버전 확인, 패치/비활성화 적용, 웹쉘 점검 — 인프라·네트워크·메일보안",
     "단기(1~2주 이내): 외부 노출 장비 패치 SLA(KEV 등재 시 72시간 이내) 규정화 — 정보보호부·인프라",
     "중장기(1분기 이내): 외부 접점 장비의 단일 벤더 의존도 평가 및 관리망 분리, EOL 장비 교체 예산 확보"],
    None, None, None,
    ["[T1] CISA / Critical Zero-Day Vulnerabilities Exploited in Citrix NetScaler ADC, Gateway / 2026-09-27 / https://www.cisa.gov/news-events/alerts/2026/09/27/critical-zero-day-vulnerabilities-exploited-citrix-netscaler-adc-gateway",
     "[T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog / 2026-09-30 / https://www.cisa.gov/news-events/alerts/2026/09/30/cisa-adds-one-known-exploited-vulnerability-catalog",
     "[T1] CISA / CISA Adds One Known Exploited Vulnerability to Catalog / 2026-10-01 / https://www.cisa.gov/news-events/alerts/2026/10/01/cisa-adds-one-known-exploited-vulnerability-catalog",
     "[T2] Rapid7 / Zero-Day Exploitation of Citrix NetScaler ADC and Gateway / 2026-09 / https://www.rapid7.com/blog/post/etr-zero-day-exploitation-of-citrix-netscaler-adc-and-gateway-cve-2026-88771-and-cve-2026-88772/",
     "[T2] Qualys / Cisco Catalyst SD-WAN Manager Authentication Bypass (CVE-2026-76504) / 2026-09-30 / https://threatprotect.qualys.com/2026/09/30/cisco-catalyst-sd-wan-manager-authentication-bypass-vulnerability-exploited-in-attacks-cve-2026-76504/"],
)

report(
    "[보고] Bitget 3.875억 달러 탈취 사고 분석 및 대응 리포트",
    ["발견 시점: 2026-09-25 탈취 실행 확인(Bitget·SlowMist·Mandiant 조사), 원인 공개 2026-10-01~10-02",
     "유출·피해 규모: 3.875억 달러(핫·웜 월렛)",
     "피해 항목: 가상자산(7개 체인). 고객 개인정보 유출 여부 미공개"],
    ["공격 유형: T1190, T1505.003, T1021, T1565 (제3자 보안장비 제로데이 2종 악용)",
     "발견 배경: 탈취 후 SlowMist·Mandiant 포렌식으로 최초 접근 2026-08-31 이상 확인",
     "발생 추정 시점: 2026-08-31(최초 접근, 조사 보도 기준)"],
    ["즉시 대응: 포렌식 및 회수 포상(5%) 프로그램 가동",
     "보안 강화: 미공개",
     "유관기관 신고: 미공개. 회수 전망은 CEO가 비관적으로 언급"],
    ["법적·규제: 전자금융거래법, 가상자산 관련 법령상 안전성·보고 의무(당사 직접 적용은 조문 확인 필요)",
     "평판·고객: 자산 보관·승인 체계에 대한 신뢰 훼손",
     "재무·운영: 직접 금전 손실, 회수 비용"],
    ["즉시(24시간 이내): 보안장비(VPN·WAF·메일·원격접속)에서 승인·정산 서버로 이어지는 접근경로 점검 — 인프라보안팀",
     "단기(1~2주 이내): 고액·비정상 승인 요청 이중 검증 및 승인 데이터 무결성 검증 적용 점검 — 승인·정산 시스템 담당",
     "중장기(1분기 이내): 제3자 보안제품 공급망 리스크 관리 및 내부 이동 탐지(EDR·NDR) 고도화"],
    None, None, None,
    ["[T3] The Hacker News / Bitget hacked via zero-day in third-party security products / 2026-10 / https://x.com/TheHackersNews/status/2105528856048496964",
     "[T4] CNBC / Bitget 'not expecting to recover a lot' from $388 million hack / 2026-10-02 / https://www.cnbc.com/2026/10/02/bitget-crypto-stolen-hack-recovery.html",
     "[T4] Forkast / The Bitget Breach: When Security Layers Become Attack Surfaces / 2026-10 / https://forkast.news/the-bitget-breach-when-security-layers-become-attack-surfaces/"],
)

# ───────────── 자체검증 ─────────────
doc.add_heading("자체검증 결과", 1)
box("[자체검증 결과]", [
    "① 논리적 비약 여부: 없음. 단, 심각도는 금융회사 영향도를 고려해 적용함 — KEV 등재 취약점은 외부 접점 장비(Citrix·FortiMail·Cisco)만 CRITICAL로, 제품 성격상 금융 영향이 불명확한 Zammad는 HIGH로 분류했고 미 국방부 DMDC 유출은 금융권 직접 영향이 낮아 MEDIUM으로 분류함(기준 문언상 CRITICAL 해석 가능).",
    "② 추측성 표현 여부: 없음. 사건 개요·메타데이터에는 확인 사실만 기재. Bitget의 북한 귀속은 CEO 주장이며 독립 검증 미확인이라 귀속하지 않음. 국내 사고의 \"중국어 오픈소스 AI 도구\"는 언론 보도 기준이며 공격 주체는 미공개로 기재.",
    "③ 실무 적용 가능성: 적합. 다만 일부 원문(CISA·금융위·FSC·BleepingComputer 등) 직접 열람이 네트워크 제한으로 차단되어 검색 결과 요약·보도에 의존함. 패치 버전·고정 릴리스·법령 조문은 벤더 공지·법령 원문 확인 필요.",
    "④ 출처 검증 상태: ✅ 확인 9건 / ⚠️ 미확인 2건 (Zammad KEV, Frontline Education)",
    "⑤ 날짜·기간 기준 검증: KST 기준 최근 7일(2026-09-27 ~ 2026-10-04) 충족. 예외: Bitget 탈취 실행(2026-09-25)과 Pentagon DMDC 침해 기간(2025-10~2026-07)은 발생 시점이 기간 외이나 공개일이 기간 내로 포함. Cisco Secure Email Gateway(2026-09-14)와 ShinyHunters PeopleSoft(2026-09-25)는 기간 외로 제외. Tier 1 중 KISA·금융보안원·개인정보보호위원회·NIST·ENISA는 검색 범위 내 해당 기간 특이사항 없음(미확인).",
], color="F0F0F0")
hr()
para("침해사고대응팀 · 생성형 AI 자동 수집·분석 · 2026-10-04", size=8.5, space=1)
para("출처: Tier 1 CISA / Tier 2 Rapid7, Qualys, runZero, Mandiant(보도 인용), Symantec(보도 인용) / Tier 3 The Hacker News, 보안뉴스, Help Net Security, SecurityAffairs, The Record, BankInfoSecurity / Tier 4 한국일보, 뉴스핌, 서울신문, 비즈워치, 이투데이, 파이낸셜뉴스, CNBC", size=8)

# Section 2 supplementary source list for items not in Section 4
doc.add_heading("부록 — SECTION 2 비상세 항목 출처", 2)
for l in [
    "Zammad KEV: [T1] CISA / CISA Adds Two Known Exploited Vulnerabilities to Catalog / 2026-10-02 / https://www.cisa.gov/news-events/alerts/2026/10/02/cisa-adds-two-known-exploited-vulnerabilities-catalog",
    "Warlock: [T2] Symantec / Warlock Ransomware Attackers Hit Water and Telecom Operators / 2026-10 / https://www.security.com/threat-intelligence/warlock-ransomware-critical-infrastructure ; [T3] The Record / https://therecord.media/warlock-ransomware-used-in-critical-infrastructure-attacks",
    "KillSec: [T3] The Record / Police disrupt KillSec ransomware, arrest suspected teenage leader / 2026-10 / https://therecord.media/killsec-ransomware-raas-arrests-europe ; [T3] BankInfoSecurity / https://www.bankinfosecurity.com/spain-arrests-teen-suspected-running-killsec-ransomware-a-33002",
    "Pentagon DMDC: [T3] CybersecurityNews / Pentagon Data Breach / 2026-10-01 / https://cybersecuritynews.com/pentagon-data-breach/ ; [T4] SAN / https://san.com/cc/pentagon-is-latest-federal-agency-hit-by-hackers/",
    "Frontline Education: [T3] BleepingComputer / Frontline Education breach exposes school district employee data / 2026-10 / https://www.bleepingcomputer.com/news/security/frontline-education-data-breach-impacts-school-district-employees/",
]:
    bullet(l)

out = "/home/user/Study/reports/주간_정보보호_동향_보고서_2026-10-04.docx"
doc.save(out)
print(out)
